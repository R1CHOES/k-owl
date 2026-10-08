import sys
import os
import json
import traceback
import re

def process_pdf(file_path):
    import pymupdf4llm
    md_text = pymupdf4llm.to_markdown(file_path)
    return {"ok": True, "markdown": md_text}

def process_docx(file_path):
    import docx
    doc = docx.Document(file_path)
    md_text = ""
    for element in doc.element.body:
        if element.tag.endswith('p'):
            for p in doc.paragraphs:
                if p._element == element:
                    text = p.text.strip()
                    if text:
                        # If it's a short title-like string, make it a heading
                        if len(text) < 80 and (text.isupper() or re.match(r'^(\d+(\.\d+)*\.?|[A-Z]\.|Section\s+\d+:?)\s+', text, re.IGNORECASE)):
                            md_text += f"## {text}\n\n"
                        else:
                            md_text += text + "\n\n"
                    break
        elif element.tag.endswith('tbl'):
            for t in doc.tables:
                if t._element == element:
                    for i, row in enumerate(t.rows):
                        row_data = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
                        md_text += "| " + " | ".join(row_data) + " |\n"
                        if i == 0:
                            md_text += "|" + "|".join(["---"] * len(row.cells)) + "|\n"
                    md_text += "\n\n"
                    break
    return {"ok": True, "markdown": md_text.strip()}

def main():
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "No file path provided."}))
        sys.exit(1)
        
    file_path = sys.argv[1]
    ext = os.path.splitext(file_path)[1].lower()
    
    try:
        if ext == '.pdf':
            result = process_pdf(file_path)
        elif ext in ['.docx', '.doc']:
            result = process_docx(file_path)
        else:
            result = {"ok": False, "error": f"Unsupported file type: {ext}"}
            
        print(json.dumps(result))
    except Exception as e:
        error_msg = f"Extraction failed: {str(e)}\n{traceback.format_exc()}"
        print(json.dumps({"ok": False, "error": error_msg}))
        sys.exit(1)

if __name__ == '__main__':
    main()
